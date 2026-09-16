"""
Core analysis engine for the client upload portal.

Given a file (PDF, JPG/JPEG/PNG, ZIP, PPTX, XLSX), returns a list of "rows"
describing each page / slide / sheet / image found, its estimated physical
size in mm, whether that matches a standard paper/slide size, and a flag
for anything unusual that needs a human look before quoting.

Design notes / honest limitations (see README for the full explanation):
  - JPG/PNG: physical size is only as good as the DPI metadata embedded in
    the file. Where no DPI tag exists we assume 300 DPI and say so.
  - PPTX: a .pptx has ONE slide size for the whole deck (PowerPoint does not
    support mixed slide sizes), so all slides share one measurement.
  - XLSX: Excel doesn't reliably store a physical "page size" unless the
    print setup was explicitly configured. Where it was, we use it
    (authoritative). Where it wasn't, we estimate the used range's physical
    size from column widths / row heights, which is an approximation, not
    an exact print-page size.
  - ZIP: unpacked recursively (depth-limited) and every supported file
    inside is analysed and listed individually, prefixed with its path
    inside the archive.
"""

import os
import zipfile
import tempfile
import shutil
from dataclasses import dataclass, field, asdict

from paper_sizes import (
    STANDARD_PAGE_SIZES,
    STANDARD_SLIDE_SIZES,
    EXCEL_PAPER_SIZE_CODES,
    LABEL_ALIASES,
    match_size,
    match_page_size,
    mm_from_points,
    mm_from_emu,
    mm_from_pixels,
)

SUPPORTED_EXTENSIONS = {".pdf", ".jpg", ".jpeg", ".png", ".zip", ".pptx", ".xlsx", ".xlsm"}
ASSUMED_IMAGE_DPI = 300
MAX_ZIP_DEPTH = 3
MAX_ZIP_UNCOMPRESSED_BYTES = 2 * 1024 * 1024 * 1024  # 2GB guard against zip bombs


@dataclass
class Row:
    source_file: str        # top-level uploaded filename
    location: str           # path within archive, or "" if not from a zip
    file_type: str          # PDF / JPG / PPTX / XLSX / ...
    unit_label: str         # "Page 1", "Slide 3", "Sheet: Budget", "Image"
    width_mm: float = None
    height_mm: float = None
    matched_size: str = ""
    is_standard: bool = True
    flagged: bool = False
    notes: str = ""

    def as_dict(self):
        d = asdict(self)
        for k in ("width_mm", "height_mm"):
            if d[k] is not None:
                d[k] = round(d[k], 1)
        return d


def analyze_pdf(path, source_file, location=""):
    rows = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                rows.append(Row(source_file, location, "PDF", "Whole document",
                                 flagged=True, notes="Password-protected - could not open to inspect pages."))
                return rows
        for i, page in enumerate(reader.pages, start=1):
            box = page.mediabox
            w_mm = mm_from_points(float(box.width))
            h_mm = mm_from_points(float(box.height))
            m = match_page_size(w_mm, h_mm)
            rows.append(Row(
                source_file, location, "PDF", f"Page {i}",
                width_mm=w_mm, height_mm=h_mm,
                matched_size=m.label, is_standard=m.is_standard,
                flagged=not m.is_standard,
                notes="" if m.is_standard else "Non-standard page size for a PDF page."
            ))
    except Exception as e:
        rows.append(Row(source_file, location, "PDF", "Whole document",
                         flagged=True, notes=f"Could not read PDF: {e}"))
    return rows


def analyze_image(path, source_file, location=""):
    rows = []
    try:
        from PIL import Image
        with Image.open(path) as img:
            width_px, height_px = img.size
            dpi = img.info.get("dpi")
            assumed = False
            if dpi and dpi[0] and dpi[1]:
                dpi_x, dpi_y = dpi
            else:
                dpi_x = dpi_y = ASSUMED_IMAGE_DPI
                assumed = True
            w_mm = mm_from_pixels(width_px, dpi_x)
            h_mm = mm_from_pixels(height_px, dpi_y)
            m = match_page_size(w_mm, h_mm)
            notes = []
            if assumed:
                notes.append(f"No DPI metadata found - assumed {ASSUMED_IMAGE_DPI} DPI to estimate print size.")
            if not m.is_standard:
                notes.append("Estimated print size doesn't match a standard paper size.")
            rows.append(Row(
                source_file, location, "Image", f"{width_px}x{height_px}px",
                width_mm=w_mm, height_mm=h_mm,
                matched_size=m.label, is_standard=m.is_standard,
                flagged=assumed or (not m.is_standard),
                notes=" ".join(notes)
            ))
    except Exception as e:
        rows.append(Row(source_file, location, "Image", "Whole file",
                         flagged=True, notes=f"Could not read image: {e}"))
    return rows


def analyze_pptx(path, source_file, location=""):
    rows = []
    try:
        from pptx import Presentation
        prs = Presentation(path)
        w_mm = mm_from_emu(prs.slide_width)
        h_mm = mm_from_emu(prs.slide_height)
        m = match_size(w_mm, h_mm, STANDARD_SLIDE_SIZES)
        slide_count = len(prs.slides)
        rows.append(Row(
            source_file, location, "PPTX", f"{slide_count} slide(s) @ deck size",
            width_mm=w_mm, height_mm=h_mm,
            matched_size=m.label, is_standard=m.is_standard,
            flagged=not m.is_standard,
            notes="" if m.is_standard else "Non-standard slide size for the whole deck."
        ))
    except Exception as e:
        rows.append(Row(source_file, location, "PPTX", "Whole file",
                         flagged=True, notes=f"Could not read PPTX: {e}"))
    return rows


def _excel_col_width_to_mm(width_chars):
    # Approximation for the default Calibri 11 font (~7px max digit width @ 96 DPI).
    if width_chars is None:
        width_chars = 8.43  # Excel default
    width_px = width_chars * 7 + 5
    return mm_from_pixels(width_px, 96)


def _excel_row_height_to_mm(height_pt):
    if height_pt is None:
        height_pt = 15.0  # Excel default row height in points
    width_px = height_pt * 96 / 72
    return mm_from_pixels(width_px, 96)


def analyze_xlsx(path, source_file, location=""):
    rows = []
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            paper_code = ws.page_setup.paperSize
            orientation = ws.page_setup.orientation
            if paper_code is not None:
                try:
                    code_int = int(paper_code)
                except (TypeError, ValueError):
                    code_int = None
                declared = EXCEL_PAPER_SIZE_CODES.get(code_int)
                if declared:
                    std_w, std_h = STANDARD_PAGE_SIZES.get(declared.split(" ")[0], (None, None))
                    display_label = LABEL_ALIASES.get(declared, declared)
                    rows.append(Row(
                        source_file, location, "Excel", f"Sheet '{ws.title}' (declared print setup)",
                        width_mm=std_w, height_mm=std_h,
                        matched_size=display_label, is_standard=True,
                        flagged=False,
                        notes=f"Orientation: {orientation or 'not set'}."
                    ))
                    continue
                else:
                    rows.append(Row(
                        source_file, location, "Excel", f"Sheet '{ws.title}' (declared print setup)",
                        matched_size=f"Unrecognised paper code {code_int}", is_standard=False,
                        flagged=True, notes="Sheet specifies a paper size we don't recognise - check manually."
                    ))
                    continue

            # No explicit paper size set - estimate physical size of the used range instead.
            dim = ws.calculate_dimension()
            if dim == "A1:A1" and ws["A1"].value is None:
                rows.append(Row(source_file, location, "Excel", f"Sheet '{ws.title}'",
                                 notes="Sheet appears empty.", flagged=False))
                continue

            min_col, min_row, max_col, max_row = openpyxl.utils.cell.range_boundaries(dim)
            total_w_mm = 0.0
            for c in range(min_col, max_col + 1):
                letter = openpyxl.utils.get_column_letter(c)
                cd = ws.column_dimensions.get(letter)
                total_w_mm += _excel_col_width_to_mm(cd.width if cd else None)
            total_h_mm = 0.0
            for r in range(min_row, max_row + 1):
                rd = ws.row_dimensions.get(r)
                total_h_mm += _excel_row_height_to_mm(rd.height if rd else None)

            m = match_page_size(total_w_mm, total_h_mm)
            rows.append(Row(
                source_file, location, "Excel", f"Sheet '{ws.title}' (estimated from used range {dim})",
                width_mm=total_w_mm, height_mm=total_h_mm,
                matched_size=m.label, is_standard=m.is_standard,
                flagged=True,  # always flag estimated (non-declared) sheets for a manual glance
                notes=("No print area/paper size set in the file - size is an ESTIMATE from column/row "
                       "dimensions, not a true print size. Please sanity-check before quoting.")
            ))
    except Exception as e:
        rows.append(Row(source_file, location, "Excel", "Whole file",
                         flagged=True, notes=f"Could not read spreadsheet: {e}"))
    return rows


DISPATCH = {
    ".pdf": analyze_pdf,
    ".jpg": analyze_image,
    ".jpeg": analyze_image,
    ".png": analyze_image,
    ".pptx": analyze_pptx,
    ".xlsx": analyze_xlsx,
    ".xlsm": analyze_xlsx,
}


def analyze_zip(path, source_file, location="", depth=0):
    rows = []
    if depth >= MAX_ZIP_DEPTH:
        rows.append(Row(source_file, location, "ZIP", "Whole archive",
                         flagged=True, notes="Archive nested too deeply - stopped unpacking for safety."))
        return rows
    tmpdir = tempfile.mkdtemp(prefix="zipscan_")
    try:
        with zipfile.ZipFile(path) as zf:
            total_size = sum(i.file_size for i in zf.infolist())
            if total_size > MAX_ZIP_UNCOMPRESSED_BYTES:
                rows.append(Row(source_file, location, "ZIP", "Whole archive",
                                 flagged=True, notes="Archive too large when uncompressed - not scanned."))
                return rows
            zf.extractall(tmpdir)
            for info in zf.infolist():
                if info.is_dir():
                    continue
                name = info.filename
                base = os.path.basename(name)
                if base.startswith(".") or "__MACOSX" in name:
                    continue
                ext = os.path.splitext(base)[1].lower()
                full_path = os.path.join(tmpdir, name)
                inner_location = f"{location}{os.path.basename(path)}/{name}" if not location else f"{location}{name}"
                if ext == ".zip":
                    rows.extend(analyze_zip(full_path, source_file, location=inner_location + " > ", depth=depth + 1))
                elif ext in DISPATCH:
                    rows.extend(DISPATCH[ext](full_path, source_file, location=inner_location))
                else:
                    rows.append(Row(source_file, inner_location, ext.lstrip(".").upper() or "Unknown",
                                     "Whole file", flagged=False,
                                     notes="File type not analysed by this tool (listed for reference only)."))
    except Exception as e:
        rows.append(Row(source_file, location, "ZIP", "Whole archive",
                         flagged=True, notes=f"Could not open archive: {e}"))
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
    return rows


def analyze_file(path, original_filename):
    """Entry point: dispatch on extension. Returns a list of Row objects."""
    ext = os.path.splitext(original_filename)[1].lower()
    if ext == ".zip":
        return analyze_zip(path, original_filename)
    elif ext in DISPATCH:
        return DISPATCH[ext](path, original_filename)
    else:
        return [Row(original_filename, "", ext.lstrip(".").upper() or "Unknown", "Whole file",
                     flagged=True, notes="Unsupported file type - not analysed.")]
