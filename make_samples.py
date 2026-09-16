"""Generate sample test files covering standard + unusual sizes for testing."""
import os
from pypdf import PdfWriter
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Emu
import openpyxl
import zipfile

OUT = os.path.join(os.path.dirname(__file__), "samples")
os.makedirs(OUT, exist_ok=True)

# --- PDF: 2 A4 pages + 1 unusual custom-size page (mixed sizes in one doc) ---
mm_to_pt = lambda mm: mm / 25.4 * 72
writer = PdfWriter()
writer.add_blank_page(width=mm_to_pt(210), height=mm_to_pt(297))  # A4
writer.add_blank_page(width=mm_to_pt(210), height=mm_to_pt(297))  # A4
writer.add_blank_page(width=mm_to_pt(150), height=mm_to_pt(400))  # unusual banner-ish size
with open(os.path.join(OUT, "brochure_mixed_sizes.pdf"), "wb") as f:
    writer.write(f)

# --- PDF: single A3 poster ---
writer2 = PdfWriter()
writer2.add_blank_page(width=mm_to_pt(420), height=mm_to_pt(297))  # A3 landscape
with open(os.path.join(OUT, "poster_a3.pdf"), "wb") as f:
    writer2.write(f)

# --- JPG: standard photo with no DPI tag (forces "assumed DPI" flag) ---
img1 = Image.new("RGB", (1748, 2480), color=(200, 200, 200))  # ~A4 at 210 DPI, but no dpi tag saved
img1.save(os.path.join(OUT, "scan_no_dpi.jpg"), "JPEG")

# --- JPG: with DPI tag matching A4 at 300dpi (2480x3508 px) ---
img2 = Image.new("RGB", (2480, 3508), color=(220, 220, 220))
img2.save(os.path.join(OUT, "scan_a4_300dpi.jpg"), "JPEG", dpi=(300, 300))

# --- PPTX: standard 16:9 widescreen, 3 slides ---
prs = Presentation()
prs.slide_width = Emu(12192000)   # 13.333in
prs.slide_height = Emu(6858000)   # 7.5in
layout = prs.slide_layouts[6]
for _ in range(3):
    prs.slides.add_slide(layout)
prs.save(os.path.join(OUT, "pitch_deck_widescreen.pptx"))

# --- PPTX: unusual custom size ---
prs2 = Presentation()
prs2.slide_width = Inches(20)
prs2.slide_height = Inches(4)
prs2.slides.add_slide(prs2.slide_layouts[6])
prs2.save(os.path.join(OUT, "banner_slide_unusual.pptx"))

# --- XLSX: no print setup (triggers estimate path) ---
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Budget"
for r in range(1, 30):
    for c in range(1, 8):
        ws.cell(row=r, column=c, value=r * c)
wb.save(os.path.join(OUT, "budget_no_pagesetup.xlsx"))

# --- XLSX: explicit A4 print setup ---
wb2 = openpyxl.Workbook()
ws2 = wb2.active
ws2.title = "Invoice"
ws2.page_setup.paperSize = ws2.PAPERSIZE_A4
ws2.page_setup.orientation = "portrait"
for r in range(1, 10):
    ws2.cell(row=r, column=1, value=f"Line {r}")
wb2.save(os.path.join(OUT, "invoice_a4.xlsx"))

# --- ZIP bundling a few of the above, plus a nested subfolder ---
zip_path = os.path.join(OUT, "client_job_bundle.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    zf.write(os.path.join(OUT, "poster_a3.pdf"), "poster_a3.pdf")
    zf.write(os.path.join(OUT, "scan_no_dpi.jpg"), "scans/scan_no_dpi.jpg")
    zf.write(os.path.join(OUT, "invoice_a4.xlsx"), "docs/invoice_a4.xlsx")

print("Sample files created in", OUT)
for f in sorted(os.listdir(OUT)):
    print(" -", f)
